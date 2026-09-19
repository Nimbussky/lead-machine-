import requests
import re
import time
import random
from bs4 import BeautifulSoup
from urllib.parse import quote_plus
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HEADERS_LIST = [
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/119.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 Chrome/120.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:121.0) Gecko/20100101 Firefox/121.0",
]

def get_headers():
    return {
        "User-Agent": random.choice(HEADERS_LIST),
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.9",
    }

PHONE_REGEX = re.compile(r'[\+]?[(]?[0-9]{1,4}[)]?[-\s\./0-9]{7,15}')
EMAIL_REGEX = re.compile(r'[a-zA-Z0-9._%+\-]+@[a-zA-Z0-9.\-]+\.[a-zA-Z]{2,}')

class Lead:
    def __init__(self):
        self.name = ""
        self.business = ""
        self.category = ""
        self.location = ""
        self.phone = ""
        self.email = ""
        self.website = ""
        self.social_links = []
        self.source = ""
        self.description = ""

    def to_dict(self):
        return {
            "Name": self.name,
            "Business": self.business,
            "Category": self.category,
            "Location": self.location,
            "Phone": self.phone,
            "Email": self.email,
            "Website": self.website,
            "Social Links": ", ".join(self.social_links) if self.social_links else "",
            "Source": self.source,
            "Description": self.description[:200],
        }

class LeadEngine:
    def __init__(self):
        self.leads = []
        self.session = requests.Session()
        self.session.headers.update(get_headers())

    def _fetch(self, url, retries=3):
        for i in range(retries):
            try:
                self.session.headers.update(get_headers())
                resp = self.session.get(url, timeout=15)
                if resp.status_code == 200:
                    return resp
                time.sleep(random.uniform(1, 3))
            except Exception:
                time.sleep(random.uniform(2, 5))
        return None

    def search_google(self, query, num_pages=3):
        leads = []
        for page in range(num_pages):
            start = page * 10
            url = f"https://www.google.com/search?q={quote_plus(query)}&start={start}"
            resp = self._fetch(url)
            if not resp:
                continue
            soup = BeautifulSoup(resp.text, "lxml")
            for g in soup.select("div.g, div[data-sokoban-container]"):
                try:
                    title_el = g.select_one("h3")
                    link_el = g.select_one("a[href]")
                    snippet_el = g.select_one("div[data-sncf], span.aCOpRe, div.VwiC3b")
                    if not link_el:
                        continue
                    lead = Lead()
                    lead.name = title_el.get_text(strip=True) if title_el else ""
                    lead.website = link_el["href"]
                    if lead.website.startswith("/url?q="):
                        lead.website = lead.website.split("/url?q=")[1].split("&")[0]
                    lead.description = snippet_el.get_text(strip=True) if snippet_el else ""
                    lead.source = "Google Search"
                    leads.append(lead)
                except Exception:
                    continue
            time.sleep(random.uniform(2, 4))
        return leads

    def search_bing(self, query, num_pages=2):
        leads = []
        for page in range(num_pages):
            first = page * 10 + 1
            url = f"https://www.bing.com/search?q={quote_plus(query)}&first={first}"
            resp = self._fetch(url)
            if not resp:
                continue
            soup = BeautifulSoup(resp.text, "lxml")
            for li in soup.select("li.b_algo"):
                try:
                    title_el = li.select_one("h2 a")
                    snippet_el = li.select_one("p, div.b_caption p")
                    if not title_el:
                        continue
                    lead = Lead()
                    lead.name = title_el.get_text(strip=True)
                    lead.website = title_el.get("href", "")
                    lead.description = snippet_el.get_text(strip=True) if snippet_el else ""
                    lead.source = "Bing Search"
                    leads.append(lead)
                except Exception:
                    continue
            time.sleep(random.uniform(2, 4))
        return leads

    def search_duckduckgo(self, query):
        leads = []
        url = f"https://html.duckduckgo.com/html/?q={quote_plus(query)}"
        resp = self._fetch(url)
        if not resp:
            return leads
        soup = BeautifulSoup(resp.text, "lxml")
        for result in soup.select("div.result"):
            try:
                title_el = result.select_one("a.result__a")
                snippet_el = result.select_one("a.result__snippet")
                if not title_el:
                    continue
                lead = Lead()
                lead.name = title_el.get_text(strip=True)
                lead.website = title_el.get("href", "")
                lead.description = snippet_el.get_text(strip=True) if snippet_el else ""
                lead.source = "DuckDuckGo"
                leads.append(lead)
            except Exception:
                continue
        return leads

    def extract_contact_from_page(self, url):
        result = {"phone": "", "email": "", "social_links": []}
        resp = self._fetch(url)
        if not resp:
            return result
        text = resp.text
        phones = PHONE_REGEX.findall(text)
        emails = EMAIL_REGEX.findall(text)
        result["phone"] = phones[0] if phones else ""
        result["email"] = emails[0] if emails else ""
        social_patterns = {
            "facebook": r'https?://(?:www\.)?facebook\.com/[^\s"\'<>]+',
            "instagram": r'https?://(?:www\.)?instagram\.com/[^\s"\'<>]+',
            "linkedin": r'https?://(?:www\.)?linkedin\.com/(?:in|company)/[^\s"\'<>]+',
            "twitter": r'https?://(?:www\.)?(?:twitter|x)\.com/[^\s"\'<>]+',
            "youtube": r'https?://(?:www\.)?youtube\.com/(?:channel|c|@)[^\s"\'<>]+',
            "tiktok": r'https?://(?:www\.)?tiktok\.com/@[^\s"\'<>]+',
            "reddit": r'https?://(?:www\.)?reddit\.com/user/[^\s"\'<>]+',
        }
        for platform, pattern in social_patterns.items():
            matches = re.findall(pattern, text)
            result["social_links"].extend(matches[:2])
        return result

    def enrich_leads(self, leads, max_enrich=15):
        enriched = 0
        for lead in leads[:max_enrich]:
            if lead.website and lead.website.startswith("http"):
                contact = self.extract_contact_from_page(lead.website)
                lead.phone = contact["phone"]
                lead.email = contact["email"]
                lead.social_links = contact["social_links"]
                enriched += 1
                time.sleep(random.uniform(1, 3))

    def search_social_platforms(self, query, location=""):
        leads = []
        search_queries = [
            f'site:linkedin.com/in "{query}" "{location}"' if location else f'site:linkedin.com/in "{query}"',
            f'site:facebook.com "{query}" "{location}"' if location else f'site:facebook.com "{query}"',
            f'site:instagram.com "{query}" "{location}"' if location else f'site:instagram.com "{query}"',
            f'site:twitter.com "{query}" OR site:x.com "{query}"',
            f'site:reddit.com "{query}" business OR portfolio',
        ]
        for sq in search_queries:
            page_leads = self.search_google(sq, num_pages=1)
            for lead in page_leads:
                if "linkedin.com" in lead.website:
                    lead.category = "LinkedIn"
                elif "facebook.com" in lead.website:
                    lead.category = "Facebook"
                elif "instagram.com" in lead.website:
                    lead.category = "Instagram"
                elif "twitter.com" in lead.website or "x.com" in lead.website:
                    lead.category = "Twitter/X"
                elif "reddit.com" in lead.website:
                    lead.category = "Reddit"
                leads.append(lead)
            time.sleep(random.uniform(2, 4))
        return leads

    def search_business_directories(self, query, location=""):
        leads = []
        dir_queries = [
            f'"{query}" "{location}" contact email phone' if location else f'"{query}" contact email phone',
            f'"{query}" business directory "{location}"' if location else f'"{query}" business directory',
            f'"{query}" freelancer OR agency OR company contact',
        ]
        for dq in dir_queries:
            page_leads = self.search_google(dq, num_pages=2)
            leads.extend(page_leads)
            time.sleep(random.uniform(2, 4))
        return leads

    def generate_leads(self, query, location="", sources=None, max_results=50):
        if sources is None:
            sources = ["google", "bing", "duckduckgo", "social", "directories"]
        all_leads = []
        if "google" in sources:
            search_q = f'{query} "{location}" contact' if location else f'{query} contact email'
            all_leads.extend(self.search_google(search_q, num_pages=3))
        if "bing" in sources:
            search_q = f'{query} "{location}"' if location else query
            all_leads.extend(self.search_bing(search_q, num_pages=2))
        if "duckduckgo" in sources:
            all_leads.extend(self.search_duckduckgo(f'{query} business contact'))
        if "social" in sources:
            all_leads.extend(self.search_social_platforms(query, location))
        if "directories" in sources:
            all_leads.extend(self.search_business_directories(query, location))
        seen = set()
        unique_leads = []
        for lead in all_leads:
            key = (lead.name.lower().strip(), lead.website.lower().strip())
            if key not in seen and lead.name:
                seen.add(key)
                lead.location = location
                lead.business = lead.business or query
                lead.category = lead.category or query
                unique_leads.append(lead)
        unique_leads = unique_leads[:max_results]
        self.enrich_leads(unique_leads)
        self.leads = unique_leads
        return unique_leads

    def export_to_excel(self, filepath):
        wb = Workbook()
        ws = wb.active
        ws.title = "AI Generated Leads"
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )
        headers = ["#", "Name", "Business", "Category", "Location", "Phone", "Email", "Website", "Social Links", "Source", "Description"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        for idx, lead in enumerate(self.leads, 1):
            data = [idx] + [lead.to_dict()[h] for h in headers[1:]]
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=idx + 1, column=col, value=str(value))
                cell.border = thin_border
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        col_widths = [5, 30, 25, 15, 20, 18, 30, 40, 50, 15, 40]
        for i, w in enumerate(col_widths):
            ws.column_dimensions[chr(65 + i)].width = w
        wb.save(filepath)
        return filepath

    def export_to_csv(self, filepath):
        import csv
        with open(filepath, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=["Name", "Business", "Category", "Location", "Phone", "Email", "Website", "Social Links", "Source", "Description"])
            writer.writeheader()
            for lead in self.leads:
                writer.writerow(lead.to_dict())
        return filepath
